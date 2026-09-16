"""文档解析器 — 支持 PDF / DOCX / TXT / MD / 代码文件"""
import io
import logging
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


async def parse_file(filename: str, content: bytes) -> str:
    """
    根据文件类型解析内容为纯文本

    支持:
        - .pdf   → pdfplumber 提取文本
        - .docx  → python-docx 提取段落
        - .txt   → 直接读取
        - .md    → 直接读取（保留 Markdown 格式）
        - .py/.js/.json/.csv/.xml/.yaml/.yml/.toml → 直接读取
        - .png/.jpg/.jpeg → 返回描述（由 Vision API 处理）
    """
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        return await _parse_pdf(content, filename)
    elif ext == ".docx":
        return await _parse_docx(content, filename)
    elif ext in (".txt", ".md", ".markdown"):
        return content.decode("utf-8", errors="replace")
    elif ext in CODE_EXTENSIONS:
        return content.decode("utf-8", errors="replace")
    elif ext in IMAGE_EXTENSIONS:
        return f"[图片文件: {filename} ({len(content)} bytes) — 图片内容由 AI 视觉模型在对话时解析]"
    else:
        # 尝试作为文本读取
        try:
            return content.decode("utf-8", errors="replace")
        except Exception:
            raise ValueError(f"不支持的文件类型: {ext}")


# 支持的文件扩展名
CODE_EXTENSIONS = {
    ".py", ".js", ".ts", ".jsx", ".tsx", ".json", ".csv", ".xml",
    ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf",
    ".html", ".css", ".scss", ".less",
    ".sql", ".sh", ".bash", ".bat", ".ps1",
    ".java", ".c", ".cpp", ".h", ".hpp", ".cs",
    ".go", ".rs", ".swift", ".kt", ".rb", ".php",
    ".r", ".R", ".lua", ".dart", ".scala", ".ex", ".exs",
}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".svg"}


async def _parse_pdf(content: bytes, filename: str) -> str:
    """解析 PDF 文件"""
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            pages = []
            for i, page in enumerate(pdf.pages):
                text = page.extract_text()
                if text and text.strip():
                    pages.append(f"--- 第 {i + 1} 页 ---\n{text.strip()}")
            if not pages:
                logger.warning(f"PDF {filename}: 未提取到文本（可能是扫描件）")
                return f"[PDF 文件: {filename} — 未提取到文本内容，可能是扫描件/图片 PDF]"
            return "\n\n".join(pages)
    except ImportError:
        logger.error("未安装 pdfplumber: pip install pdfplumber")
        return f"[错误] PDF 解析需要安装 pdfplumber: pip install pdfplumber"
    except Exception as e:
        logger.exception(f"PDF 解析失败: {filename}")
        return f"[PDF 解析错误] {filename}: {e}"


async def _parse_docx(content: bytes, filename: str) -> str:
    """解析 Word 文档"""
    try:
        import docx
        doc = docx.Document(io.BytesIO(content))
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            # 处理标题
            if para.style and para.style.name.startswith("Heading"):
                level = para.style.name.replace("Heading", "").strip() or "1"
                paragraphs.append(f"{'#' * int(level)} {text}")
            else:
                paragraphs.append(text)

        if not paragraphs:
            return f"[Word 文件: {filename} — 未提取到文本内容]"

        # 也提取表格
        for table in doc.tables:
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(" | ".join(cells))
            if rows:
                paragraphs.append("\n表格:\n" + "\n".join(rows))

        return "\n\n".join(paragraphs)
    except ImportError:
        return f"[错误] Word 解析需要安装 python-docx: pip install python-docx"
    except Exception as e:
        logger.exception(f"DOCX 解析失败: {filename}")
        return f"[DOCX 解析错误] {filename}: {e}"


def get_file_type(filename: str) -> str:
    """获取文件类型分类"""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return "pdf"
    elif ext == ".docx":
        return "word"
    elif ext in CODE_EXTENSIONS:
        return "code"
    elif ext in IMAGE_EXTENSIONS:
        return "image"
    else:
        return "text"
